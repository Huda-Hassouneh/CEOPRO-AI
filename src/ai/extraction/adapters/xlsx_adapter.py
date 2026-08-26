"""
CEOPRO AI - XLSX Source Adapter.
Converts the first row of a worksheet into headers and every subsequent
non-empty row into a dict. Requires openpyxl (pip install openpyxl).
"""
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook


def read_xlsx_file(
    file_path: str, sheet_name: Optional[str] = None
) -> Tuple[List[str], List[Dict[str, object]]]:
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    sheet = workbook[sheet_name] if sheet_name else workbook.active

    rows_iter = sheet.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if header_row is None:
        return [], []

    headers = [str(h).strip() if h is not None else "" for h in header_row]
    rows: List[Dict[str, object]] = []
    for raw_row in rows_iter:
        if all(cell is None for cell in raw_row):
            continue
        row = {headers[i]: raw_row[i] for i in range(len(headers)) if i < len(raw_row)}
        rows.append(row)

    return headers, rows
