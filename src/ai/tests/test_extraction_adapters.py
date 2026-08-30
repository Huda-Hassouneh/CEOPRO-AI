"""
Unit tests for extraction/adapters/*.py. Zero coverage anywhere in the
repo before this file - found during a QA pass that also found these
files had no caller at all (see test_extraction_file_dispatch.py).
"""
import os
import tempfile

from openpyxl import Workbook

from src.ai.extraction.adapters.api_adapter import flatten_records, read_paginated_api
from src.ai.extraction.adapters.csv_adapter import read_csv_file
from src.ai.extraction.adapters.db_adapter import read_db_query
from src.ai.extraction.adapters.pdf_adapter import read_pdf_file
from src.ai.extraction.adapters.xlsx_adapter import read_xlsx_file


def _write_temp(suffix: str) -> str:
    return os.path.join(tempfile.gettempdir(), f"adapter_test{suffix}")


def test_read_csv_file_parses_headers_and_rows():
    path = _write_temp(".csv")
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        f.write("product_name,quantity\nWidget A,5\nWidget B,3\n")
    headers, rows = read_csv_file(path)
    assert headers == ["product_name", "quantity"]
    assert rows == [
        {"product_name": "Widget A", "quantity": "5"},
        {"product_name": "Widget B", "quantity": "3"},
    ]


def test_read_csv_file_strips_byte_order_mark():
    path = _write_temp("_bom.csv")
    with open(path, "wb") as f:
        f.write(b"\xef\xbb\xbfproduct_name,quantity\r\nWidget,5\r\n")
    headers, rows = read_csv_file(path)
    assert headers[0] == "product_name"  # not "﻿product_name"


def test_read_xlsx_file_parses_first_sheet():
    path = _write_temp(".xlsx")
    wb = Workbook()
    ws = wb.active
    ws.append(["product_name", "quantity"])
    ws.append(["Widget A", 5])
    ws.append(["Widget B", 3])
    wb.save(path)

    headers, rows = read_xlsx_file(path)
    assert headers == ["product_name", "quantity"]
    assert rows == [
        {"product_name": "Widget A", "quantity": 5},
        {"product_name": "Widget B", "quantity": 3},
    ]


def test_read_xlsx_file_skips_fully_blank_rows():
    path = _write_temp("_blank.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.append(["product_name", "quantity"])
    ws.append(["Widget A", 5])
    ws.append([None, None])
    ws.append(["Widget B", 3])
    wb.save(path)

    headers, rows = read_xlsx_file(path)
    assert len(rows) == 2  # the blank row must not appear


def test_read_xlsx_file_empty_sheet_returns_empty():
    path = _write_temp("_empty.xlsx")
    wb = Workbook()
    wb.save(path)
    headers, rows = read_xlsx_file(path)
    assert headers == []
    assert rows == []


class _FakePage:
    """
    Shared fake for both PDF test paths below - extract_tables() defaults
    to no tables found (the common "unstructured PDF" case) so passing
    only `text` reproduces the pre-upgrade fallback-only test unchanged;
    passing `tables` exercises the new structured-extraction-first path.
    Mocks pdfplumber.open rather than generating a real PDF (no PDF-
    writing library is a project dependency) - this tests read_pdf_file()'s
    own logic independent of pdfplumber's own extraction correctness,
    which isn't this project's code to test.
    """

    def __init__(self, text=None, tables=None):
        self._text = text
        self._tables = tables or []

    def extract_text(self):
        return self._text

    def extract_tables(self):
        return self._tables


def _fake_pdf(pages):
    class FakePdf:
        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

    FakePdf.pages = pages
    return FakePdf()


def test_read_pdf_file_extracts_text_per_page_when_no_table_found(monkeypatch):
    """Regression: unstructured PDFs (no detectable table on any page)
    must keep behaving exactly as before the structured-extraction
    upgrade - one row per page, blank/None-text pages skipped."""
    import src.ai.extraction.adapters.pdf_adapter as pdf_adapter_module

    pages = [
        _FakePage(text="Invoice total 45.00 JOD"),
        _FakePage(text="   "),
        _FakePage(text=None),
    ]
    monkeypatch.setattr(pdf_adapter_module.pdfplumber, "open", lambda path: _fake_pdf(pages))

    headers, rows = read_pdf_file("fake.pdf")
    assert headers == ["page_text"]
    assert len(rows) == 1  # the blank and None-text pages must be skipped
    assert "45.00" in rows[0]["page_text"]


def test_read_pdf_file_extracts_a_real_table_as_structured_rows(monkeypatch):
    """The core upgrade: a PDF whose page contains an actual table (the
    canonical template's own header row) is parsed as (headers, rows) -
    typed/structured, not dumped as a raw text blob - so it can reach
    TEMPLATE_COMPLIANT the same way a CSV/XLSX upload does."""
    import src.ai.extraction.adapters.pdf_adapter as pdf_adapter_module

    table = [
        ["product_name", "quantity", "unit_price"],
        ["Widget A", "5", "19.99"],
        ["Widget B", "2", "9.99"],
    ]
    pages = [_FakePage(tables=[table])]
    monkeypatch.setattr(pdf_adapter_module.pdfplumber, "open", lambda path: _fake_pdf(pages))

    headers, rows = read_pdf_file("fake.pdf")
    assert headers == ["product_name", "quantity", "unit_price"]
    assert rows == [
        {"product_name": "Widget A", "quantity": "5", "unit_price": "19.99"},
        {"product_name": "Widget B", "quantity": "2", "unit_price": "9.99"},
    ]


def test_read_pdf_file_table_spanning_multiple_pages_shares_one_header(monkeypatch):
    """A standard multi-page form: the header row repeats identically on
    each page's table (a common real-world PDF export pattern) - rows
    from every page accumulate under the one shared header set."""
    import src.ai.extraction.adapters.pdf_adapter as pdf_adapter_module

    header = ["product_name", "quantity", "unit_price"]
    pages = [
        _FakePage(tables=[[header, ["Widget A", "5", "19.99"]]]),
        _FakePage(tables=[[header, ["Widget B", "2", "9.99"]]]),
    ]
    monkeypatch.setattr(pdf_adapter_module.pdfplumber, "open", lambda path: _fake_pdf(pages))

    headers, rows = read_pdf_file("fake.pdf")
    assert headers == header
    assert len(rows) == 2
    assert rows[0]["product_name"] == "Widget A"
    assert rows[1]["product_name"] == "Widget B"


def test_read_pdf_file_skips_a_table_whose_header_does_not_match_the_first_one(monkeypatch):
    """A later page's table with a genuinely different shape is skipped,
    not guessed at by position - misaligning cells under the wrong header
    would be a worse outcome than not using that table at all."""
    import src.ai.extraction.adapters.pdf_adapter as pdf_adapter_module

    pages = [
        _FakePage(tables=[[["product_name", "quantity", "unit_price"], ["Widget A", "5", "19.99"]]]),
        _FakePage(tables=[[["unrelated_col_a", "unrelated_col_b"], ["x", "y"]]]),
    ]
    monkeypatch.setattr(pdf_adapter_module.pdfplumber, "open", lambda path: _fake_pdf(pages))

    headers, rows = read_pdf_file("fake.pdf")
    assert headers == ["product_name", "quantity", "unit_price"]
    assert len(rows) == 1
    assert rows[0]["product_name"] == "Widget A"


def test_read_pdf_file_ignores_single_row_table_false_positives(monkeypatch):
    """pdfplumber's table detector can return a single-row false positive
    from an unrelated bordered box/rule on a page - a "table" with no data
    rows at all isn't usable and must not become the document's canonical
    header set, especially if a real table follows on a later page."""
    import src.ai.extraction.adapters.pdf_adapter as pdf_adapter_module

    pages = [
        _FakePage(tables=[[["Just One Row"]]]),
        _FakePage(tables=[[["product_name", "quantity", "unit_price"], ["Widget A", "5", "19.99"]]]),
    ]
    monkeypatch.setattr(pdf_adapter_module.pdfplumber, "open", lambda path: _fake_pdf(pages))

    headers, rows = read_pdf_file("fake.pdf")
    assert headers == ["product_name", "quantity", "unit_price"]
    assert len(rows) == 1


def test_read_pdf_file_handles_none_header_cells_from_merged_columns(monkeypatch):
    """extract_tables() can legally return None for a blank/merged header
    cell (confirmed against the installed pdfplumber's own type hints) -
    must not crash trying to .strip() it."""
    import src.ai.extraction.adapters.pdf_adapter as pdf_adapter_module

    table = [
        ["product_name", None, "quantity", "unit_price"],
        ["Widget A", None, "5", "19.99"],
    ]
    pages = [_FakePage(tables=[table])]
    monkeypatch.setattr(pdf_adapter_module.pdfplumber, "open", lambda path: _fake_pdf(pages))

    headers, rows = read_pdf_file("fake.pdf")
    assert headers[0] == "product_name"
    assert headers[1] == ""  # normalized from None, not a crash


def test_read_db_query_maps_column_names_to_dicts():
    class FakeCursor:
        description = [("product_id",), ("quantity",)]

        def execute(self, query, params):
            pass

        def fetchall(self):
            return [("p1", 5), ("p2", 3)]

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

    class FakeConn:
        def cursor(self):
            return FakeCursor()

    headers, rows = read_db_query(FakeConn(), "SELECT product_id, quantity FROM x;")
    assert headers == ["product_id", "quantity"]
    assert rows == [{"product_id": "p1", "quantity": 5}, {"product_id": "p2", "quantity": 3}]


def test_flatten_records_unions_headers_across_records_preserving_order():
    records = [{"a": 1, "b": 2}, {"b": 3, "c": 4}]
    headers, rows = flatten_records(records)
    assert headers == ["a", "b", "c"]
    assert rows == records


def test_read_paginated_api_stops_when_a_page_returns_empty():
    pages = {1: [{"id": 1}], 2: [{"id": 2}], 3: []}

    def fetch_page(page_num):
        return pages.get(page_num)

    headers, rows = read_paginated_api(fetch_page)
    assert headers == ["id"]
    assert rows == [{"id": 1}, {"id": 2}]


def test_read_paginated_api_respects_max_pages():
    calls = []

    def fetch_page(page_num):
        calls.append(page_num)
        return [{"id": page_num}]  # never returns empty - would loop forever without max_pages

    read_paginated_api(fetch_page, max_pages=3)
    assert calls == [1, 2, 3]
