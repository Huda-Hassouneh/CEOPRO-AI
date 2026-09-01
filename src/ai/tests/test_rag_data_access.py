"""
Offline tests for data_access.py's pure logic: the pgvector text-literal
helpers, and fetch_document_text()'s extension-based extraction dispatch
(2026-09-01 - PDF/DOCX/XLSX support). No psycopg2 adapter is used for the
`vector` type (see the module docstring), so the literal round-trip
(Python array -> Postgres literal -> Python array) is the actual mechanism
the persistence fix (audit finding P1) depends on being correct.
DB-touching functions (replace_document_chunks, load_tenant_chunks,
list_documents) are covered by test_rag_integration.py's live-DB tests
instead - these tests only need a fake MinIO client, no real Postgres.

PDF extraction is tested against a mocked pdfplumber.open(), matching
test_extraction_adapters.py's established convention (no PDF-writing
library is a project dependency, so a real PDF can't be generated in-test)
- this tests fetch_document_text()'s own dispatch logic, not pdfplumber's
extraction correctness, which isn't this project's code. DOCX and XLSX
extraction are each tested against a real, in-memory file generated with
python-docx/openpyxl themselves (both are writers as well as readers),
which PDF has no equivalent for - more rigorous where the tooling allows it.
"""

import io

import numpy as np
import pytest

import src.ai.rag.data_access as data_access_module
from src.ai.rag.data_access import UnsupportedDocumentTypeError, _from_pgvector_literal, _to_pgvector_literal, fetch_document_text


def test_round_trips_a_realistic_384_dim_embedding():
    original = np.random.RandomState(42).normal(size=384).astype(np.float32)
    literal = _to_pgvector_literal(original)
    restored = _from_pgvector_literal(literal)

    assert restored.shape == original.shape
    np.testing.assert_allclose(restored, original, rtol=1e-5)


def test_literal_format_matches_pgvector_bracket_syntax():
    literal = _to_pgvector_literal(np.array([0.1, -0.2, 0.3], dtype=np.float32))
    assert literal.startswith("[")
    assert literal.endswith("]")
    assert literal.count(",") == 2  # 3 values -> 2 separators


def test_negative_and_zero_values_round_trip_correctly():
    original = np.array([0.0, -1.5, 1.5], dtype=np.float32)
    restored = _from_pgvector_literal(_to_pgvector_literal(original))
    np.testing.assert_allclose(restored, original, rtol=1e-6)


class _FakeMinioResponse:
    def __init__(self, data: bytes):
        self._data = data

    def read(self):
        return self._data

    def close(self):
        pass

    def release_conn(self):
        pass


class _FakeMinioClient:
    def __init__(self, data: bytes):
        self._data = data

    def get_object(self, bucket, object_key):
        return _FakeMinioResponse(self._data)


def test_fetch_document_text_reads_plain_txt():
    client = _FakeMinioClient("Sunscreen SPF 50 is our best seller.".encode("utf-8"))
    assert fetch_document_text(client, "bucket", "policy.txt") == "Sunscreen SPF 50 is our best seller."


def test_fetch_document_text_reads_markdown_as_plain_text():
    client = _FakeMinioClient("# Policy\nReturns within 30 days.".encode("utf-8"))
    assert fetch_document_text(client, "bucket", "policy.md") == "# Policy\nReturns within 30 days."


def test_fetch_document_text_raises_for_an_unsupported_extension():
    client = _FakeMinioClient(b"whatever bytes")
    with pytest.raises(UnsupportedDocumentTypeError, match="\\.pptx"):
        fetch_document_text(client, "bucket", "slides.pptx")


def test_fetch_document_text_raises_for_no_extension_at_all():
    client = _FakeMinioClient(b"whatever bytes")
    with pytest.raises(UnsupportedDocumentTypeError, match="no extension"):
        fetch_document_text(client, "bucket", "README")


class _FakePdfPage:
    def __init__(self, text):
        self._text = text

    def extract_text(self):
        return self._text


class _FakePdf:
    def __init__(self, pages):
        self.pages = pages

    def __enter__(self):
        return self

    def __exit__(self, *args):
        return False


def test_fetch_document_text_extracts_pdf_text_per_page(monkeypatch):
    """Mocks pdfplumber.open() - see this file's own module docstring for why."""
    pages = [_FakePdfPage("Page one content."), _FakePdfPage("Page two content.")]
    monkeypatch.setattr(data_access_module.pdfplumber, "open", lambda file_obj: _FakePdf(pages))

    client = _FakeMinioClient(b"%PDF-1.4 fake pdf bytes")
    text = fetch_document_text(client, "bucket", "policy.pdf")

    assert "Page one content." in text
    assert "Page two content." in text


def test_fetch_document_text_skips_pdf_pages_with_no_extractable_text(monkeypatch):
    """A scanned/image-only page returns None from extract_text() - must be
    dropped, not turned into a literal "None" string in the joined output."""
    pages = [_FakePdfPage("Real content."), _FakePdfPage(None)]
    monkeypatch.setattr(data_access_module.pdfplumber, "open", lambda file_obj: _FakePdf(pages))

    client = _FakeMinioClient(b"%PDF-1.4 fake pdf bytes")
    text = fetch_document_text(client, "bucket", "policy.pdf")

    assert text == "Real content."


def test_fetch_document_text_extracts_docx_paragraphs():
    """Real .docx, generated in-memory with python-docx itself (a writer as
    well as a reader) - no mocking needed, unlike the PDF case above."""
    import docx

    document = docx.Document()
    document.add_paragraph("Our return policy allows returns within 30 days.")
    document.add_paragraph("Sunscreen SPF 50 is our best selling product.")
    buffer = io.BytesIO()
    document.save(buffer)

    client = _FakeMinioClient(buffer.getvalue())
    text = fetch_document_text(client, "bucket", "policy.docx")

    assert "Our return policy allows returns within 30 days." in text
    assert "Sunscreen SPF 50 is our best selling product." in text


def test_fetch_document_text_skips_empty_docx_paragraphs():
    import docx

    document = docx.Document()
    document.add_paragraph("Real content.")
    document.add_paragraph("")  # a blank line in the source document
    document.add_paragraph("   ")  # whitespace-only
    buffer = io.BytesIO()
    document.save(buffer)

    client = _FakeMinioClient(buffer.getvalue())
    text = fetch_document_text(client, "bucket", "policy.docx")

    assert text == "Real content."


def test_fetch_document_text_extracts_xlsx_rows_tab_separated():
    """Real .xlsx, generated in-memory with openpyxl itself - no mocking
    needed, same reasoning as the DOCX tests above."""
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append(["Product", "Price", "Currency"])
    sheet.append(["Premium Olive Oil 1L", 24.5, "JOD"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    client = _FakeMinioClient(buffer.getvalue())
    text = fetch_document_text(client, "bucket", "catalog.xlsx")

    assert "Sheet: Products" in text
    assert "Product\tPrice\tCurrency" in text
    assert "Premium Olive Oil 1L\t24.5\tJOD" in text


def test_fetch_document_text_covers_every_sheet_in_a_multi_sheet_xlsx():
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet1 = workbook.active
    sheet1.title = "Products"
    sheet1.append(["Sunscreen SPF 50"])
    sheet2 = workbook.create_sheet("Policies")
    sheet2.append(["Returns accepted within 30 days"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    client = _FakeMinioClient(buffer.getvalue())
    text = fetch_document_text(client, "bucket", "handbook.xlsx")

    assert "Sheet: Products" in text
    assert "Sunscreen SPF 50" in text
    assert "Sheet: Policies" in text
    assert "Returns accepted within 30 days" in text


def test_fetch_document_text_skips_fully_blank_xlsx_rows_and_renders_empty_cells_as_empty():
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Name", "Notes"])
    sheet.append(["Widget A", None])  # a genuinely blank cell, not the string "None"
    sheet.append([None, None])  # a fully blank row - must be dropped, not rendered as an empty line
    buffer = io.BytesIO()
    workbook.save(buffer)

    client = _FakeMinioClient(buffer.getvalue())
    text = fetch_document_text(client, "bucket", "notes.xlsx")

    assert "None" not in text
    assert "Widget A\t" in text
    lines = [line for line in text.splitlines() if line.strip()]
    assert len(lines) == 3  # "Sheet: ..." + header row + "Widget A" row, no blank-row artifact
