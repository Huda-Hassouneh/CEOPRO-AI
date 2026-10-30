"""Local meeting fixture: real public APIs, one reserved synthetic identity.

Run with the checkout Python: tools/dev/seed_meeting_demo.py prepare|seed|verify.
Private state, source inputs, backup and evidence stay under .dev. No provider
credentials, administrator connection or entitlement override enter the seeder.
"""
from __future__ import annotations

import argparse
from contextlib import closing
from dataclasses import replace
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
import hashlib
import io
import json
from pathlib import Path
import re
import secrets
import shutil
import subprocess
import sys
import time
import uuid

import dev

sys.path.insert(0, str(dev.ROOT / "backend"))
import database
import httpx
from openpyxl import Workbook
from psycopg2 import sql
from scripts.demo_companies.profiles import ZAYTOUNA
from scripts.demo_companies.seed import Client, configure_company
from scripts.ci_smoke_account import _complete_guided_onboarding

EMAIL = "meeting-phase1@keel-demo.example"
BUSINESS = "Zaytouna Meeting Demo — SYNTHETIC"
VERSION = "meeting-phase1-v1"
END = date(2026, 10, 27)
DAYS = 180
BASE = dev.STATE / "fixtures" / "meeting-phase1"
EVIDENCE = dev.STATE / "evidence" / "meeting-phase1"
ACCOUNT = BASE / "account.private.json"
JOURNAL = BASE / "journal.json"
COMPANY = replace(ZAYTOUNA, business_name=BUSINESS, owner_name="Synthetic Demo Owner",
                  target_markets=[{"market_type": "COUNTRY", "country_code": "JO"}],
                  products=ZAYTOUNA.products[:10])
PROTECTED = ("companies", "tenant_users", "products", "inventory", "invoices", "invoice_items",
             "transactions", "tenant_competitors", "competitor_prices", "billing_subscriptions",
             "billing_payments", "billing_trials", "tenant_plan_assignments", "tenant_credit_balances")


def write(path, value):
    dev.write_json(path, value)


def read(path, default=None):
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else default


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def fingerprints(conn, *, tenants=None, tables=None):
    result = {}
    with conn.cursor() as cur:
        cur.execute("SELECT table_name FROM information_schema.tables WHERE table_schema='public' AND table_type='BASE TABLE' ORDER BY table_name")
        present = [r[0] for r in cur.fetchall()]
        for table in (tables or present):
            if table not in present:
                continue
            where = sql.SQL("")
            args = ()
            if tenants is not None:
                cur.execute("SELECT 1 FROM information_schema.columns WHERE table_schema='public' AND table_name=%s AND column_name='tenant_id'", (table,))
                if not cur.fetchone():
                    continue
                where = sql.SQL("WHERE tenant_id::text = ANY(%s)")
                args = (tenants,)
            cur.execute(sql.SQL("SELECT count(*), md5(coalesce(string_agg(h,'' ORDER BY h),'')) FROM (SELECT md5(row_to_json(t)::text) h FROM public.{} t {}) s").format(sql.Identifier(table), where), args)
            result[table] = list(cur.fetchone())
    return result


def prepare(c):
    """Export one consistent snapshot; restore and compare every public table."""
    marker = EVIDENCE / "backup.json"
    if marker.exists():
        saved = read(marker)
        if saved["database"] != "keel_dev_" + c["id"]:
            raise RuntimeError("Backup belongs to another database")
        if digest(Path(saved["dump"])) != saved["dump_sha256"] or saved["status"] != "PASS":
            raise RuntimeError("Verified backup is missing or changed")
        print("PASS: existing verified backup retained")
        return
    dev.check_docker()
    dev.verify_containers(c)
    database.check(c)
    container = dev.docker(c, "ps", "-q", "postgres", capture=True).strip()
    if not container or "\n" in container:
        raise RuntimeError("Expected one owned PostgreSQL container")
    target = dev.STATE / "backups" / (VERSION + "-" + datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ"))
    target.mkdir(parents=True, exist_ok=False)
    protected = {}
    for rel in (".dev/config.json", ".dev/integrations.env", "backend/.env.example"):
        source = dev.ROOT / rel
        if source.exists():
            protected[rel] = digest(source)
            dest = target / "private-config" / rel
            dest.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source, dest)
    name = "keel_dev_" + c["id"]
    restored = "keel_dev_demo_restore_" + uuid.uuid4().hex[:12]
    with closing(database.connect(c, external=False)) as conn:
        conn.set_session(isolation_level="REPEATABLE READ", readonly=True)
        with conn.cursor() as cur:
            cur.execute("SELECT pg_export_snapshot()")
            snapshot = cur.fetchone()[0]
            cur.execute("SELECT tenant_id::text FROM companies ORDER BY tenant_id")
            tenants = [r[0] for r in cur.fetchall()]
        before = fingerprints(conn)
        customer = fingerprints(conn, tenants=tenants, tables=PROTECTED)
        with (target / "database.dump").open("wb") as out, (target / "backup.private.log").open("wb") as log:
            subprocess.run(["docker", "exec", container, "pg_dump", "-U", "keel_dev_admin", "-d", name,
                            "-Fc", "--snapshot=" + snapshot], stdout=out, stderr=log, check=True)
    with (target / "roles.private.sql").open("wb") as out, (target / "roles.private.log").open("wb") as log:
        subprocess.run(["docker", "exec", container, "pg_dumpall", "-U", "keel_dev_admin", "--roles-only"], stdout=out, stderr=log, check=True)
    with closing(database.connect(c, external=False)) as conn:
        conn.autocommit = True
        with conn.cursor() as cur:
            cur.execute(sql.SQL("CREATE DATABASE {} TEMPLATE template0").format(sql.Identifier(restored)))
    with (target / "database.dump").open("rb") as inp, (target / "restore.private.log").open("wb") as log:
        subprocess.run(["docker", "exec", "-i", container, "pg_restore", "-U", "keel_dev_admin", "-d", restored,
                        "--exit-on-error"], stdin=inp, stdout=log, stderr=log, check=True)
    with closing(database.connect(c, database=restored, external=False)) as conn:
        if fingerprints(conn) != before:
            raise RuntimeError("Restored table fingerprints differ from exported snapshot")
    result = {"status": "PASS", "database": name, "dump": str(target / "database.dump"),
              "dump_sha256": digest(target / "database.dump"), "restore_database": restored,
              "tables_verified": len(before), "existing_tenants": tenants,
              "customer_fingerprints": customer, "protected_hashes": protected}
    write(target / "manifest.json", result)
    write(marker, result)
    print(f"PASS: backup restored; {len(before)} table fingerprints match; retained {restored}")


def inputs():
    """Fixed dates, Decimal money and explicit stock snapshots; no AI outputs."""
    BASE.mkdir(parents=True, exist_ok=True)
    products = []
    closing_stock = [140, 8, 0, 25, 320, 4.5, 180, 150, 95, 70]
    for i, p in enumerate(COMPANY.products):
        products.append({"sku": f"DEMO-{i+1:03}", "name": p.name_en, "ar": p.name_local,
                         "category": p.category, "price": str(Decimal(str(p.price_history[-1][1]))),
                         "cost": str(Decimal(str(p.cost))), "stock": str(closing_stock[i]),
                         "reorder": "10", "safety": "5"})
    lines, invoices = [], []
    for day_index in range(DAYS):
        day = END - timedelta(days=DAYS - 1 - day_index)
        for group in range(2):
            number = f"DEMO-{day:%Y%m%d}-{group+1}"
            subtotal = Decimal(0)
            for i in range(group * 5, group * 5 + 5):
                p = products[i]
                # Modest daily POS summaries, one fractional weight SKU; weekly
                # and monthly shape provides real history for existing models.
                quantity = Decimal(2 + (day_index + i * 3) % 5 + (2 if day.weekday() in (4, 5) else 0)
                                   + day_index // 60)
                if i == 5:
                    quantity /= 4
                total = quantity * Decimal(p["price"])
                subtotal += total
                lines.append({"date": day.isoformat(), "invoice": number, "sku": p["sku"],
                              "quantity": str(quantity), "unit_price": p["price"], "total": str(total),
                              "gross_profit": str(total - quantity * Decimal(p["cost"]))})
            invoices.append({"number": number, "date": day.isoformat(), "subtotal": str(subtotal)})
    manifest = {"version": VERSION, "synthetic": True, "email": EMAIL, "business": BUSINESS,
                "currency": "JOD", "timezone": "Asia/Amman", "start": invoices[0]["date"], "end": END.isoformat(),
                "products": products, "expected_counts": {"products": 10, "inventory": 10, "invoices": 360,
                "invoice_items": 1800, "transactions": 1800, "competitors": 3, "competitor_prices": 24,
                "business_import_jobs": 3, "knowledge_documents": 3},
                "totals": {"revenue": str(sum(Decimal(x["total"]) for x in lines)),
                           "units": str(sum(Decimal(x["quantity"]) for x in lines)),
                           "gross_profit": str(sum(Decimal(x["gross_profit"]) for x in lines))},
                "invoices": invoices, "lines": lines}
    write(BASE / "truth-inputs.json", manifest)
    docs = {
        "company-and-delivery.md": """# Zaytouna Meeting Demo — SYNTHETIC business input
Amman specialty grocery. Currency JOD, timezone Asia/Amman.
This fictional business sells olive oil, spices and dates. All demo names,
customers, prices and sales are synthetic; competitor observations are manual
scenario inputs, never scraped evidence. Delivery is inside Amman only.
Delivery charge: JOD 2.000; free delivery for basket subtotal JOD 30.000 or more.
Order cutoff 14:00 Asia/Amman for next-business-day delivery. Friday is closed.
""",
        "returns-and-stock.md": """# Zaytouna Meeting Demo — SYNTHETIC returns and stock policy
Unopened packaged goods may be returned within 7 calendar days with receipt.
Opened food is not returnable unless damaged or incorrect. Report damaged goods
within 48 hours. Refunds use the original payment method after inspection.
The stock file is a closing physical count on 2026-10-27, after historical sales
and replenishments. It is not an opening balance and historical sales must not
be subtracted again. All quantities are at Amman Demo Store; loose za'atar is kg.
No purchase/replenishment events are asserted by this snapshot.
""",
        "harvest-marketing-brief.md": """# Zaytouna Meeting Demo — SYNTHETIC marketing input
Campaign brief only; no generated campaign or image exists yet.
Product: Extra Virgin Olive Oil 1L (DEMO-001). Retail price JOD 9.500.
Goal: promote the autumn pantry range to shoppers in Amman, in English or Arabic.
Tone: warm and factual. Channel: Instagram. Do not claim organic certification,
health benefits, a discount, free delivery below JOD 30.000, or unsupported awards.
Suggested Phase 2 request: create one harvest-themed post using the actual product.
"""}
    for name, content in docs.items():
        (BASE / name).write_text(content, encoding="utf-8")
    return manifest


def client_for(c):
    class SeedClient(Client):
        def request(self, method, path, expected=(200,), **kwargs):
            headers = {**self.headers, **kwargs.pop("headers", {})}
            response = self.http.request(method, path, headers=headers, **kwargs)
            if response.status_code not in expected:
                raise RuntimeError(f"API {method} {path} failed (HTTP {response.status_code}).")
            return response.json() if response.content else None
    return SeedClient(f"http://127.0.0.1:{c['ports']['api']}")


def login(c, *, create=False):
    account = read(ACCOUNT)
    if not account:
        if not create:
            raise RuntimeError("Run seed to create the designated demo account")
        account = {"email": EMAIL, "business": BUSINESS, "password": "Demo!" + secrets.token_urlsafe(24)}
        write(ACCOUNT, account)
    if account["email"] != EMAIL or account["business"] != BUSINESS:
        raise RuntimeError("Private identity file differs from designated synthetic identity")
    client = client_for(c)
    response = client.http.post("/auth/login", json={"email": EMAIL, "password": account["password"]})
    if response.status_code == 401 and create:
        if account.get("tenant_id"):
            raise RuntimeError("Existing demo login failed; refusing to create a replacement")
        if not account.get("registration_token"):
            registration = client.request("POST", "/auth/signup", expected=(202,), json={
                "email": EMAIL, "password": account["password"], "full_name": COMPANY.owner_name,
                "business_name": BUSINESS, "country_code": "JO", "primary_currency": "JOD"})
            account["registration_token"] = registration["registration_token"]
            write(ACCOUNT, account)
        # Read only the designated reserved recipient from the checkout's local
        # Mailpit inbox. Codes and challenge tokens are never logged.
        mail = httpx.Client(base_url=f"http://127.0.0.1:{c['ports']['inbox']}", timeout=20)
        code = None
        for _ in range(10):
            listing = mail.get("/api/v1/search", params={"query": "to:" + EMAIL}).json()
            messages = listing.get("messages", [])
            for message in messages:
                body = mail.get("/api/v1/message/" + message["ID"]).json()
                if not any(item.get("Address", "").casefold() == EMAIL for item in body.get("To", [])):
                    continue
                match = re.search(r"(?<!\d)\d{6}(?!\d)", body.get("Text", ""))
                if match:
                    code = match.group()
                    break
            if code:
                break
            time.sleep(1)
        mail.close()
        if not code:
            raise RuntimeError("No verification code delivered to the designated local recipient")
        client.request("POST", "/auth/email/verify", json={"email": EMAIL, "code": code,
                       "registration_token": account["registration_token"]})
        response = client.http.post("/auth/login", json={"email": EMAIL, "password": account["password"]})
    if response.status_code != 200:
        raise RuntimeError(f"Demo login failed HTTP {response.status_code}")
    session = response.json()
    if session.get("business_name") != BUSINESS or session.get("email_verified") is not True:
        raise RuntimeError("Login returned an unexpected or unverified workspace")
    if account.get("tenant_id") and account["tenant_id"] != session["tenant_id"]:
        raise RuntimeError("Demo tenant binding changed")
    backup = read(EVIDENCE / "backup.json")
    if session["tenant_id"] in backup["existing_tenants"]:
        raise RuntimeError("Refusing to seed a tenant that existed before the dedicated seed")
    account.update({k: session[k] for k in ("tenant_id", "user_id")})
    account.pop("registration_token", None)
    write(ACCOUNT, account)
    client.headers = {"Authorization": "Bearer " + session["access_token"]}
    return client, account


def workbook_bytes(invoices, lines, products):
    from src.ai.extraction.structured_business import HEADERS
    book = Workbook()
    book.remove(book.active)
    sheet = book.create_sheet("Invoices")
    sheet.append(HEADERS["Invoices"])
    for x in invoices:
        sheet.append([x["number"], "Synthetic POS daily summary", x["date"], x["date"],
                      x["subtotal"], "0", x["subtotal"], "JOD", "PAID"])
    sheet = book.create_sheet("Invoice Items")
    sheet.append(HEADERS["Invoice Items"])
    for x in lines:
        p = products[x["sku"]]
        sheet.append([x["invoice"], p["id"], p["name"], x["sku"], x["quantity"], x["unit_price"]])
    # ZIP timestamps are otherwise wall-clock-dependent. Normalize the archive
    # so payload identity stays stable when regenerating the same fixture.
    out = io.BytesIO()
    book.save(out)
    book.close()
    import zipfile
    stable = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(out.getvalue())) as source, zipfile.ZipFile(stable, "w", zipfile.ZIP_DEFLATED) as dest:
        for name in source.namelist():
            content = source.read(name)
            if name == "docProps/core.xml":
                content = re.sub(rb"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ", b"2026-10-27T00:00:00Z", content)
            item = zipfile.ZipInfo(name, (2026, 10, 27, 0, 0, 0))
            item.compress_type = zipfile.ZIP_DEFLATED
            dest.writestr(item, content)
    return stable.getvalue()


def post_import(client, kind, name, payload, journal):
    key = VERSION + ":" + name
    entry = journal.get(key)
    sha = hashlib.sha256(payload).hexdigest()
    if entry:
        if entry["sha256"] != sha:
            raise RuntimeError("Generated import payload changed; refusing a second identity")
    else:
        preview = client.request("POST", "/extraction/structured-imports/preview", params={"kind": kind},
                                 headers={**client.headers, "Idempotency-Key": key},
                                 files={"file": (name, payload)})
        if preview["issues"] or preview["counts"]["needsAttention"]:
            raise RuntimeError("Structured import has unresolved review findings")
        entry = {"sha256": sha, "draft": preview["draftId"], "reviewHash": preview["reviewHash"]}
        journal[key] = entry
        write(JOURNAL, journal)
    # Replay the exact original reviewed confirmation. Re-reviewing committed
    # invoices would correctly report an existing invoice identity conflict.
    result = client.request("POST", f"/extraction/structured-imports/{entry['draft']}/confirm",
                            json={"choices": {}, "reviewHash": entry["reviewHash"]})
    if entry.get("result") and entry["result"] != result:
        raise RuntimeError("Confirmation replay changed its response")
    entry["result"] = result
    write(JOURNAL, journal)
    print(f"PASS: {name} confirmed/replayed ({result['kind']})")


def seed(c):
    prepare(c)
    truth = inputs()
    client, account = login(c, create=True)
    journal = read(JOURNAL, {})
    replay_tables = (*PROTECTED, "data_sources", "ingestion_jobs", "rag_documents_metadata", "rag_document_chunks",
                     "tenant_usage_counters", "keel_structured_import_posts", "keel_import_batches")
    replay_before = None
    if journal.get("seed_complete"):
        with closing(database.connect(c, external=False)) as conn:
            replay_before = fingerprints(conn, tenants=[account["tenant_id"]], tables=replay_tables)
    subscription = client.request("GET", "/billing/subscription")
    if not subscription.get("hasAccess"):
        raise RuntimeError("Designated tenant has no valid signup trial; no rights will be invented")
    if "onboarded" not in journal:
        configure_company(client, COMPANY)
        _complete_guided_onboarding(client, COMPANY)
        journal["onboarded"] = True
        write(JOURNAL, journal)
    existing = client.request("GET", "/catalog/products", params={"pageSize": 100})["products"]
    mapped = {p["sku"]: p for p in existing}
    if any(p["sku"] not in {x["sku"] for x in truth["products"]} for p in existing):
        raise RuntimeError("Unexpected catalog content in designated tenant")
    for p in truth["products"]:
        if p["sku"] not in mapped:
            created = client.request("POST", "/products", json={"product_name": {"en": p["name"], "ar": p["ar"]},
                                     "sku": p["sku"], "brand": {"en": "Zaytouna Demo"},
                                     "category": {"en": p["category"]},
                                     "description": {"en": "SYNTHETIC meeting input; retail packaged goods"},
                                     "current_price": float(p["price"]), "cost_price": float(p["cost"]), "currency": "JOD"})
            mapped[p["sku"]] = {"id": created["product_id"], "name": {"en": p["name"]}}
        else:
            actual = mapped[p["sku"]]
            if actual["name"]["en"] != p["name"] or Decimal(actual["price"]) != Decimal(p["price"]) or Decimal(actual["costPrice"]) != Decimal(p["cost"]):
                raise RuntimeError("Existing product differs from source; refusing to overwrite it")
    products = {p["sku"]: {**p, "id": mapped[p["sku"]]["id"]} for p in truth["products"]}
    for index in range(2):
        invoices = truth["invoices"][index * 180:(index + 1) * 180]
        numbers = {x["number"] for x in invoices}
        lines = [x for x in truth["lines"] if x["invoice"] in numbers]
        payload = workbook_bytes(invoices, lines, products)
        name = f"synthetic-sales-{index+1}.xlsx"
        (BASE / name).write_bytes(payload)
        post_import(client, "invoices", name, payload, journal)
    import csv
    from src.ai.extraction.structured_business import HEADERS
    text = io.StringIO(newline="")
    writer = csv.writer(text)
    writer.writerow(HEADERS["Inventory"])
    for p in products.values():
        writer.writerow([p["id"], p["name"], p["sku"], p["stock"], p["reorder"], p["safety"], "Amman Demo Store"])
    payload = text.getvalue().encode("utf-8-sig")
    (BASE / "synthetic-closing-stock.csv").write_bytes(payload)
    post_import(client, "inventory", "synthetic-closing-stock.csv", payload, journal)
    tracked = {x["name"]: x["id"] for x in client.request("GET", "/competitors")["competitors"]}
    for ci in range(3):
        name = ("Demo Cedar Pantry", "Demo Amman Grocer", "Demo Harvest Foods")[ci] + " (SYNTHETIC)"
        cid = tracked.get(name)
        if not cid:
            cid = client.request("POST", "/competitors", json={"name": name,
                                 "website_url": f"https://meeting-rival-{ci+1}.example",
                                 "competition_level": "PRODUCT", "product_ids": [p["id"] for p in list(products.values())[:4]]})["competitor_id"]
        prices = client.request("GET", f"/competitors/{cid}/prices")["prices"]
        for pi, p in enumerate(list(products.values())[:4]):
            for age in (14, 1):
                price = (Decimal(p["price"]) * (Decimal("0.95") + Decimal(ci) / 20)
                         + (Decimal("0.100") if age == 14 else Decimal(0))).quantize(Decimal(".001"))
                observed = datetime.combine(END - timedelta(days=age), datetime.min.time(), timezone.utc).isoformat()
                if not any(x["product_id"] == p["id"] and Decimal(str(x["price"])) == price and
                           datetime.fromisoformat(x["observed_at"]).date().isoformat() == observed[:10] for x in prices):
                    client.request("POST", f"/competitors/{cid}/prices", json={"product_id": p["id"], "price": str(price),
                                   "currency": "JOD", "observed_at": observed,
                                   "product_url": f"https://meeting-rival-{ci+1}.example/synthetic/{p['sku']}"})
    for path in sorted(BASE.glob("*.md")):
        result = client.request("POST", "/rag/documents/upload", files={"file": (path.name, path.read_bytes(), "text/markdown")})
        journal[path.name] = {k: v for k, v in result.items() if k in ("document_id", "status", "searchable", "duplicate")}
        write(JOURNAL, journal)
        print(f"Knowledge input {path.name}: {result.get('status')}; searchable={result.get('searchable')}")
    write(BASE / "marketing-request.json", {"product_id": products["DEMO-001"]["id"], "platforms": ["INSTAGRAM"],
           "style": "Warm autumn pantry", "tone": "Warm and factual", "focus": "Amman harvest pantry",
           "request_key": VERSION + "-marketing", "generate_image": False})
    write(EVIDENCE / "seed-identity.json", {"tenant_id": account["tenant_id"], "email": EMAIL, "business": BUSINESS})
    journal["seed_complete"] = True
    write(JOURNAL, journal)
    if replay_before is not None:
        with closing(database.connect(c, external=False)) as conn:
            replay_after = fingerprints(conn, tenants=[account["tenant_id"]], tables=replay_tables)
        replay = {"status": "PASS" if replay_before == replay_after else "FAIL", "tables": len(replay_before),
                  "changed_tables": [t for t in replay_before if replay_before[t] != replay_after[t]],
                  "before": replay_before, "after": replay_after}
        write(EVIDENCE / "idempotency.json", replay)
        if replay["status"] != "PASS":
            raise RuntimeError("Replay changed business or allowance state; inspect private idempotency.json")
    print("PASS: business inputs seeded; existing automatic outputs are reported separately")
    verify(c)


def verify(c):
    truth = inputs()
    client, account = login(c)
    report = {"checked_at": datetime.now(timezone.utc).isoformat(), "commit": dev.run(["git", "rev-parse", "HEAD"], capture=True).strip(),
              "tenant_id": account["tenant_id"], "email": EMAIL, "checks": [], "feature_reads": {}}
    def check(name, expected, actual):
        report["checks"].append({"name": name, "expected": str(expected), "actual": str(actual), "status": "PASS" if expected == actual else "FAIL"})
    with closing(database.connect(c, external=False)) as conn, conn.cursor() as cur:
        for name, expected in truth["expected_counts"].items():
            table = {"competitors": "tenant_competitors", "knowledge_documents": "rag_documents_metadata",
                     "business_import_jobs": "ingestion_jobs"}.get(name, name)
            filter_sql = sql.SQL(" AND storage_bucket_path LIKE 'uploads/%%'") if name == "knowledge_documents" else sql.SQL("")
            cur.execute(sql.SQL("SELECT count(*) FROM {} WHERE tenant_id=%s{}").format(sql.Identifier(table), filter_sql), (account["tenant_id"],))
            check(name, expected, cur.fetchone()[0])
        cur.execute("SELECT coalesce(sum(total_amount),0) FROM invoices WHERE tenant_id=%s AND payment_status='PAID'", (account["tenant_id"],))
        check("total_revenue", Decimal(truth["totals"]["revenue"]), cur.fetchone()[0])
        cur.execute("SELECT sum(ii.quantity),sum(ii.total_price-ii.quantity*p.cost_price) FROM invoice_items ii JOIN products p ON p.tenant_id=ii.tenant_id AND p.product_id=ii.product_id WHERE ii.tenant_id=%s", (account["tenant_id"],))
        units, profit = cur.fetchone()
        check("total_units", Decimal(truth["totals"]["units"]), units)
        check("total_gross_profit", Decimal(truth["totals"]["gross_profit"]), profit)
        cur.execute("SELECT p.sku,ii.quantity,ii.unit_price,ii.total_price,i.invoice_number,(i.issue_date AT TIME ZONE 'Asia/Amman')::date FROM invoice_items ii JOIN products p ON p.tenant_id=ii.tenant_id AND p.product_id=ii.product_id JOIN invoices i ON i.tenant_id=ii.tenant_id AND i.invoice_id=ii.invoice_id WHERE ii.tenant_id=%s", (account["tenant_id"],))
        observed_lines = sorted((s, str(d), n, q, u, t) for s, q, u, t, n, d in cur.fetchall())
        expected_lines = sorted((x["sku"], x["date"], x["invoice"], Decimal(x["quantity"]), Decimal(x["unit_price"]), Decimal(x["total"])) for x in truth["lines"])
        check("all_invoice_line_facts", True, expected_lines == observed_lines)
        cur.execute("SELECT p.sku,i.stock_quantity,i.reorder_level,i.safety_stock FROM inventory i JOIN products p ON p.tenant_id=i.tenant_id AND p.product_id=i.product_id WHERE i.tenant_id=%s", (account["tenant_id"],))
        stock = sorted(cur.fetchall())
        expected_stock = sorted((p["sku"], Decimal(p["stock"]), Decimal(p["reorder"]), Decimal(p["safety"])) for p in truth["products"])
        check("all_stock_facts", True, stock == expected_stock)
        cur.execute("SELECT job_status,count(*) FROM ingestion_jobs WHERE tenant_id=%s GROUP BY job_status", (account["tenant_id"],))
        check("completed_import_jobs", [("COMPLETED", 3)], cur.fetchall())
        report["generated_output_counts"] = {}
        for table in ("demand_forecasts", "campaigns", "conversation_messages"):
            cur.execute(sql.SQL("SELECT count(*) FROM {} WHERE tenant_id=%s").format(sql.Identifier(table)), (account["tenant_id"],))
            report["generated_output_counts"][table] = cur.fetchone()[0]
        cur.execute("SELECT count(*) FROM rag_documents_metadata WHERE tenant_id=%s AND left(storage_bucket_path,11)='_generated/'", (account["tenant_id"],))
        report["generated_output_counts"]["application_knowledge_summaries"] = cur.fetchone()[0]
        cur.execute("SELECT count(*) FROM billing_payments WHERE tenant_id=%s", (account["tenant_id"],))
        check("Stripe_payments", 0, cur.fetchone()[0])
        backup = read(EVIDENCE / "backup.json")
        preserved = fingerprints(conn, tenants=backup["existing_tenants"], tables=PROTECTED)
        for table, old in backup["customer_fingerprints"].items():
            check("existing_tenants_unchanged:" + table, old, preserved[table])
    # Compare every tenant-owned table with the verified pre-seed restore, in
    # addition to the explicit business/billing table checks above.
    with closing(database.connect(c, database=backup["restore_database"], external=False)) as conn:
        all_before = fingerprints(conn, tenants=backup["existing_tenants"])
    with closing(database.connect(c, external=False)) as conn:
        all_after = fingerprints(conn, tenants=backup["existing_tenants"])
    check("all_existing_tenant_owned_tables_unchanged", True, all_before == all_after)
    report["preservation_table_count"] = len(all_before)
    report["preservation_changed_tables"] = [name for name in all_before if all_before[name] != all_after[name]]
    for rel, sha in backup["protected_hashes"].items():
        check("protected_file:" + rel, sha, digest(dev.ROOT / rel))
    catalog = client.request("GET", "/catalog/products", params={"pageSize": 100})
    check("API_product_count", 10, catalog["pagination"]["total"])
    check("API_stock_snapshot_count", 10, len(client.request("GET", "/forecasting/overview", params={"auto_refresh": "false"})["products"]))
    inventory = client.request("GET", "/dashboard/inventory")
    check("API_low_stock_count", 3, inventory["low_stock_count"])
    check("API_in_stock_percent", Decimal("90"), Decimal(str(inventory["in_stock_pct"])))
    now = datetime.now(timezone.utc)
    from zoneinfo import ZoneInfo
    today = now.astimezone(ZoneInfo("Asia/Amman")).date()
    report["truth_windows"] = {}
    for days in (7, 30, 90):
        # Matches core.business_time.window_start (inclusive calendar dates).
        start = today - timedelta(days=days-1)
        expected = sum((Decimal(x["total"]) for x in truth["lines"] if date.fromisoformat(x["date"]) >= start), Decimal(0))
        metrics = client.request("GET", "/dashboard/metrics", params={"window_days": days})
        check(f"dashboard_revenue_{days}d", expected, Decimal(str(metrics["revenue"]["amount"])))
        window_lines = [x for x in truth["lines"] if date.fromisoformat(x["date"]) >= start]
        expected_units = sum((Decimal(x["quantity"]) for x in window_lines), Decimal(0))
        check(f"dashboard_quantities_{days}d", expected_units, Decimal(str(metrics["units_sold"]["units"])))
        trend = client.request("GET", "/dashboard/sales-trend", params={"window_days": days})
        check(f"sales_trend_revenue_{days}d", expected, sum((Decimal(str(x["value"])) for x in trend["points"]), Decimal(0)))
        report["truth_windows"][str(days)] = {"as_of": today.isoformat(), "start_inclusive": start.isoformat(),
                                            "revenue": str(expected), "quantity_sum": str(expected_units)}
    sales = client.request("GET", "/reports/sales", params={"limit": 5000})
    check("API_report_rows", 1800, sales["total_row_count"])
    check("API_report_revenue", Decimal(truth["totals"]["revenue"]), sum((Decimal(str(x["total_price"])) for x in sales["rows"]), Decimal(0)))
    comparisons = client.request("GET", "/pricing/comparisons", params={"product_ids": ",".join(p["id"] for p in catalog["products"])})
    check("pricing_covered_products", 4, sum(bool(p["covered"]) for p in comparisons["products"].values()))
    subscription = client.request("GET", "/billing/subscription")
    check("normal_signup_trial", ("trial", "trialing", True, False),
          tuple(subscription.get(k) for k in ("source", "status", "hasAccess", "hasSubscription")))
    workspaces = client.request("GET", "/auth/workspaces")
    check("single_isolated_workspace", [account["tenant_id"]], [x["tenant_id"] for x in workspaces["workspaces"]])
    check("platform_owner_access_denied", 403, client.http.get("/platform/me", headers=client.headers).status_code)
    docs = client.request("GET", "/rag/documents")["documents"]
    report["knowledge_documents"] = [{k: x.get(k) for k in ("file_name", "status", "searchable", "chunk_count")} for x in docs]
    for name, path, params in (("inventory", "/dashboard/inventory", {}), ("data_connections", "/onboarding/sources", {}),
                               ("imports", "/extraction/imports", {}), ("demand_prediction", "/forecasting/overview", {"auto_refresh": "false"}),
                               ("pricing", "/pricing/comparisons", {"product_ids": ",".join(p["id"] for p in catalog["products"])}), ("marketing_options", "/marketing/options", {}),
                               ("marketing_suggestions", "/marketing/suggestions", {}), ("billing", "/billing/subscription", {})):
        response = client.http.get(path, headers=client.headers, params=params)
        report["feature_reads"][name] = {"http": response.status_code, "status": "PASS" if response.status_code == 200 else "BLOCKED"}
        if response.status_code == 200:
            write(EVIDENCE / (name + ".json"), response.json())
    forecast = read(EVIDENCE / "demand_prediction.json")
    report["actual_forecasts"] = [{k: p.get(k) for k in ("name", "forecast_next_30", "forecast_method", "forecast_created_at")}
                                  for p in forecast["products"] if p.get("forecast_next_30") is not None]
    # Native RLS with the actual restricted password, using the existing
    # service-principal context: even an explicit foreign tenant filter yields 0.
    from src.ai import db
    import os
    previous = os.environ.copy()
    try:
        os.environ.update(dev.environment(c, external=False))
        conn = db.app_role_connection(account["tenant_id"], account["user_id"])
        try:
            with conn.cursor() as cur:
                cur.execute("SELECT current_user")
                check("runtime_role", "ceopro_app", cur.fetchone()[0])
                cur.execute("SELECT count(*) FROM products WHERE tenant_id::text=ANY(%s)", (backup["existing_tenants"],))
                check("RLS_foreign_products_visible", 0, cur.fetchone()[0])
                cur.execute("SELECT count(*) FROM invoices WHERE tenant_id::text=ANY(%s)", (backup["existing_tenants"],))
                check("RLS_foreign_invoices_visible", 0, cur.fetchone()[0])
        finally:
            conn.close()
        # The pre-existing tenants have empty catalogs. Check the populated
        # direction too: a different tenant context cannot see this fixture.
        conn = db.app_role_connection(str(uuid.UUID(int=0)), str(uuid.UUID(int=0)))
        try:
            with conn.cursor() as cur:
                cur.execute("SELECT count(*) FROM products WHERE tenant_id=%s", (account["tenant_id"],))
                check("RLS_unrelated_context_sees_demo_products", 0, cur.fetchone()[0])
                cur.execute("SELECT count(*) FROM invoices WHERE tenant_id=%s", (account["tenant_id"],))
                check("RLS_unrelated_context_sees_demo_invoices", 0, cur.fetchone()[0])
        finally:
            conn.close()
    finally:
        os.environ.clear()
        os.environ.update(previous)
    ready = all(x.get("searchable") for x in docs) and len(docs) == 3
    report["RAG_indexing"] = "PASS" if ready else "BLOCKED"
    report["status"] = ("FAIL" if any(x["status"] == "FAIL" for x in report["checks"]) else
                        "BLOCKED" if not ready or any(x["status"] != "PASS" for x in report["feature_reads"].values()) else "PASS")
    write(EVIDENCE / "verification.json", report)
    print(f"{report['status']}: {len(report['checks'])} reconciliation/preservation checks; RAG indexing {report['RAG_indexing']}")
    for x in report["checks"]:
        if x["status"] != "PASS":
            print(f"FAIL {x['name']}: expected {x['expected']}, actual {x['actual']}")
    if report["status"] != "PASS":
        raise RuntimeError("Seed truth reconciliation failed; inspect private verification.json")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("action", choices=("prepare", "seed", "verify"))
    args = parser.parse_args()
    c = dev.config()
    if Path(sys.executable).resolve() != dev.PYTHON.resolve():
        return subprocess.call([str(dev.PYTHON), str(Path(__file__).resolve()), args.action])
    dev.private_state()
    EVIDENCE.mkdir(parents=True, exist_ok=True)
    with dev.command_lock():
        {"prepare": prepare, "seed": seed, "verify": verify}[args.action](c)
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:
        # Never stringify network/DB errors that could include credentials.
        if isinstance(exc, RuntimeError):
            print(str(exc), file=sys.stderr)
        else:
            print(f"BLOCKED: {type(exc).__name__}; no secrets printed. Inspect local runtime logs.", file=sys.stderr)
        sys.exit(1)
