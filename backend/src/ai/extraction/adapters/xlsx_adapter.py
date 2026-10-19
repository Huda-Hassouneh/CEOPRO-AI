"""
KEEL - XLSX Source Adapter.
Converts the first row of a worksheet into headers and every subsequent
non-empty row into a dict. Requires openpyxl (pip install openpyxl).
"""
from datetime import date, datetime, time
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook

from src.ai.core.office_files import InvalidOfficeFileError, ensure_safe_office_archive


def _cell_value(value):
    """Excel stores dates as real date cells: hand them on as ISO text (a
    midnight datetime is just a date), never as a datetime object nothing
    downstream can serialise or parse as text."""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    return value


def read_xlsx_file(
    file_path: str, sheet_name: Optional[str] = None
) -> Tuple[List[str], List[Dict[str, object]]]:
    ensure_safe_office_archive(file_path)  # zip-bomb / not-a-workbook check before parsing
    try:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises KeyError/OSError/InvalidFileException for broken workbooks
        raise InvalidOfficeFileError(
            "This Excel file couldn't be opened. Open it in Excel, save it again and re-upload, or export it as CSV."
        ) from exc
    sheet = workbook[sheet_name] if sheet_name else workbook.active

    rows_iter = sheet.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if header_row is None:
        return [], []

    headers = [str(h).strip() if h is not None else "" for h in header_row]
    rows: List[Dict[str, object]] = []
    for raw_row in rows_iter:
        if all(cell is None for cell in raw_row):
            continue
        row = {headers[i]: _cell_value(raw_row[i]) for i in range(len(headers)) if i < len(raw_row)}
        rows.append(row)

    return headers, rows
