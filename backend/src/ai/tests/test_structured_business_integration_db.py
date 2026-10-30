import io,json,os
from decimal import Decimal
from uuid import uuid4
from concurrent.futures import ThreadPoolExecutor
import pytest,psycopg2
from openpyxl import Workbook
from src.ai import db
from src.ai.routers.products import create_product,upsert_inventory
from src.ai.tests.test_products_integration_db import seeded_tenants
from src.ai.tests.test_durable_upload_integration_db import setup

def seed(setup):
    client,app,ts,us=setup;c=db.app_role_connection(ts[0],us[0])
    try:p=create_product(c,ts[0],us[0],{'en':'Reference','ar':'مرجع'},None,None,None,10,None,'JOD',sku='REF-1');upsert_inventory(c,ts[0],p['product_id'],Decimal('8'));c.commit();return p
    finally:c.close()
def preview(client,kind,payload,name='fixture.csv'):
    response=client.post('/extraction/structured-imports/preview',params={'kind':kind},headers={'Idempotency-Key':str(uuid4())},files={'file':(name,payload)});assert response.status_code==200,response.text;return response.json()
def review(client,p,choices=None):
    r=client.post('/extraction/structured-imports/'+p['draftId']+'/review',json={'choices':choices or {}});assert r.status_code==200,r.text;return r.json()
def confirm(client,p,r,choices=None):return client.post('/extraction/structured-imports/'+p['draftId']+'/confirm',json={'choices':choices or {},'reviewHash':r['reviewHash']})
def stock_csv(quantity=0):return f'sku,stock_quantity,reorder_level,safety_stock,warehouse_location\nREF-1,{quantity},2,1,Main\n'.encode()
def invoice_file(status='PAID'):
    from src.ai.extraction.structured_business import HEADERS
    book=Workbook();book.remove(book.active);a=book.create_sheet('Invoices');a.append(HEADERS['Invoices']);a.append(['INV-'+uuid4().hex[:8],'Synthetic Customer','2026-10-26','2026-11-01',20,2,22,'JOD',status]);i=book.create_sheet('Invoice Items');i.append(HEADERS['Invoice Items']);i.append([a.cell(2,1).value,'','','REF-1',2,10]);out=io.BytesIO();book.save(out);book.close();return out.getvalue()
def query(tenant,sql):
    with psycopg2.connect(os.environ['AI_TEST_DATABASE_URL']) as c,c.cursor() as q:q.execute(sql,(tenant,));return q.fetchall()

def test_zero_stock_explicit_review_atomic_post_replay_and_audit(setup):
    client,app,ts,us=setup;product=seed(setup);p=preview(client,'inventory',stock_csv());assert p['matches'][0]['status']=='matched'
    assert query(ts[0],'SELECT stock_quantity FROM inventory WHERE tenant_id=%s')[0][0]==8
    r=review(client,p);assert r['changes'][0]['previous']=='8.0000'
    result=confirm(client,p,r);assert result.status_code==200,result.text
    assert query(ts[0],'SELECT stock_quantity FROM inventory WHERE tenant_id=%s')[0][0]==0
    assert confirm(client,p,r).json()==result.json()
    assert len(query(ts[0],"SELECT evidence FROM keel_structured_import_posts WHERE tenant_id=%s"))==1
    assert len(query(ts[0],"SELECT changed_data_json FROM audit_logs WHERE tenant_id=%s AND action_type='STRUCTURED_IMPORT_POSTED'"))==1

@pytest.mark.parametrize('status,transactions',[('PAID',1),('UNPAID',0),('CANCELLED',0)])
def test_invoice_items_preserve_tax_status_and_forecast_transactions_without_double_sales(setup,status,transactions):
    client,app,ts,us=setup;seed(setup);p=preview(client,'invoices',invoice_file(status),'invoice.xlsx');r=review(client,p);result=confirm(client,p,r);assert result.status_code==200,result.text
    assert query(ts[0],'SELECT subtotal,tax_amount,total_amount,payment_status FROM invoices WHERE tenant_id=%s')==[(Decimal('20'),Decimal('2'),Decimal('22'),status)]
    assert len(query(ts[0],'SELECT transaction_id FROM transactions WHERE tenant_id=%s'))==transactions
    assert len(query(ts[0],'SELECT invoice_id FROM invoice_items WHERE tenant_id=%s'))==1
    assert query(ts[0],'SELECT stock_quantity FROM inventory WHERE tenant_id=%s')[0][0]==8
    assert confirm(client,p,r).json()==result.json()
    from datetime import date
    from src.ai.forecasting.data_access import load_daily_demand
    connection=db.app_role_connection(ts[0],us[0])
    try:
        with connection.cursor() as cursor:cursor.execute('SELECT product_id FROM products WHERE tenant_id=%s',(ts[0],));pid=str(cursor.fetchone()[0])
        demand=load_daily_demand(connection,ts[0],pid,through=date(2026,10,26))
        # Preserve Python demand semantics: non-cancelled invoice demand includes unpaid sales.
        assert float(demand['quantity'].sum())==(0 if status=='CANCELLED' else 2)
    finally:connection.close()

def test_stale_review_rolls_back_all_changes(setup):
    client,app,ts,us=setup;p0=seed(setup);p=preview(client,'inventory',stock_csv(5));r=review(client,p)
    c=db.app_role_connection(ts[0],us[0]);upsert_inventory(c,ts[0],p0['product_id'],Decimal(7));c.commit();c.close()
    assert confirm(client,p,r).status_code==409
    assert not query(ts[0],'SELECT batch_id FROM keel_structured_import_posts WHERE tenant_id=%s')
    assert query(ts[0],'SELECT stock_quantity FROM inventory WHERE tenant_id=%s')[0][0]==7

def test_conflicting_identifiers_and_cross_tenant_drafts_cannot_be_confirmed(setup):
    from src.ai.core.security import create_access_token
    client,app,ts,us=setup;seed(setup);payload=b'sku,product_name,stock_quantity,reorder_level,safety_stock\nREF-1,Wrong,5,2,1\n';p=preview(client,'inventory',payload);assert p['matches'][0]['status']=='conflict' and p['issues']
    r=review(client,p);assert confirm(client,p,r).status_code==409
    client.headers['Authorization']='Bearer '+create_access_token(ts[1],us[1],email_verified=True)[0]
    assert client.post('/extraction/structured-imports/'+p['draftId']+'/review',json={}).status_code==404

def test_template_endpoint_is_real_xlsx_and_unknown_headers_fail(setup):
    client,app,ts,us=setup
    assert client.get('/extraction/structured-imports/templates/invoices/xlsx').content.startswith(b'PK')
    r=client.post('/extraction/structured-imports/preview',params={'kind':'inventory'},files={'file':('bad.csv',b'stock_quantity,sku,unsafe\n1,A,2\n')});assert r.status_code==422


def test_concurrent_confirm_once_and_staff_cannot_review_or_post(setup):
    client,app,ts,us=setup;seed(setup);p=preview(client,'inventory',stock_csv(3));r=review(client,p)
    with ThreadPoolExecutor(max_workers=2) as pool:responses=list(pool.map(lambda _:confirm(client,p,r),[1,2]))
    assert [x.status_code for x in responses]==[200,200]
    assert responses[0].json()==responses[1].json()
    assert len(query(ts[0],'SELECT batch_id FROM keel_structured_import_posts WHERE tenant_id=%s'))==1
    assert query(ts[0],"SELECT sum(used) FROM tenant_usage_counters WHERE tenant_id=%s AND feature_code='document_extraction'")[0][0]==1
    with psycopg2.connect(os.environ['AI_TEST_DATABASE_URL']) as c,c.cursor() as q:q.execute("UPDATE tenant_users SET role_key='staff' WHERE tenant_id=%s AND user_id=%s",(ts[0],us[0]))
    assert client.post('/extraction/structured-imports/'+p['draftId']+'/review',json={}).status_code==403
    assert confirm(client,p,r).status_code==403

def test_create_missing_product_is_deferred_until_atomic_confirmation(setup):
    client,app,ts,us=setup;p=preview(client,'inventory',b'sku,product_name,stock_quantity,reorder_level,safety_stock\nNEW-1,New canonical,4,2,1\n');assert p['matches'][0]['status']=='missing'
    assert query(ts[0],'SELECT count(*) FROM products WHERE tenant_id=%s')[0][0]==0
    choices={p['matches'][0]['key']:{'action':'create','price':'2.5000','currency':'JOD','sku':'NEW-1'}}
    r=review(client,p,choices);result=confirm(client,p,r,choices);assert result.status_code==200,result.text
    assert query(ts[0],'SELECT sku,current_price FROM products WHERE tenant_id=%s')==[('NEW-1',Decimal('2.5'))]
    assert query(ts[0],'SELECT stock_quantity FROM inventory WHERE tenant_id=%s')[0][0]==4

def test_multiple_locations_do_not_silently_collapse_and_invalid_sku_is_422(setup):
    client,app,ts,us=setup;seed(setup);p=preview(client,'inventory',b'sku,stock_quantity,reorder_level,safety_stock,warehouse_location\nREF-1,1,0,0,A\nREF-1,2,0,0,B\n');assert p['issues'];r=review(client,p);assert confirm(client,p,r).status_code==409
    bad=client.post('/extraction/structured-imports/preview',params={'kind':'inventory'},files={'file':('bad.csv',b'sku,stock_quantity,reorder_level,safety_stock\nbad space,1,0,0\n')});assert bad.status_code==422


def test_restricted_role_post_evidence_is_append_only_and_staff_cannot_forge_it(setup):
    client,app,ts,us=setup;seed(setup);p=preview(client,'inventory',stock_csv());r=review(client,p);assert confirm(client,p,r).status_code==200
    c=db.app_role_connection(ts[0],us[0])
    try:
        with c.cursor() as q:
            q.execute("SELECT relrowsecurity,relforcerowsecurity FROM pg_class WHERE oid='keel_structured_import_posts'::regclass");assert q.fetchone()==(True,True)
            q.execute("SELECT has_table_privilege(current_user,'keel_structured_import_posts','UPDATE')");assert q.fetchone()[0] is False
            with pytest.raises(psycopg2.errors.InsufficientPrivilege):q.execute("UPDATE keel_structured_import_posts SET response='{}' WHERE tenant_id=%s",(ts[0],))
        c.rollback()
    finally:c.close()
    p2=preview(client,'inventory',stock_csv(2))
    with psycopg2.connect(os.environ['AI_TEST_DATABASE_URL']) as owner,owner.cursor() as q:q.execute("UPDATE tenant_users SET role_key='staff' WHERE tenant_id=%s AND user_id=%s",(ts[0],us[0]))
    c=db.app_role_connection(ts[0],us[0])
    try:
        with c.cursor() as q:
            with pytest.raises(psycopg2.errors.InsufficientPrivilege):q.execute("INSERT INTO keel_structured_import_posts(tenant_id,batch_id,actor_id,request_hash,response,evidence) VALUES(%s,%s,%s,%s,'{}','{}')",(ts[0],p2['draftId'],us[0],'a'*64))
        c.rollback()
    finally:c.close()


def test_company_deletion_reports_retained_reviewed_import_without_erasing_it(setup):
    from src.ai import company_data
    client,app,ts,us=setup;seed(setup);p=preview(client,'inventory',stock_csv());r=review(client,p);assert confirm(client,p,r).status_code==200
    with psycopg2.connect(os.environ['AI_TEST_DATABASE_URL']) as conn:
        with pytest.raises(company_data.DeletionBlocked,match='integration_evidence_review'):company_data.delete_company(conn,ts[0])
        conn.rollback()
        with conn.cursor() as q:q.execute('SELECT count(*) FROM keel_structured_import_posts WHERE tenant_id=%s',(ts[0],));assert q.fetchone()==(1,)
    from src.ai.routers import settings
    app.include_router(settings.router)
    r=client.request('DELETE','/settings/company',json={'confirm_name':query(ts[0],'SELECT business_name FROM companies WHERE tenant_id=%s')[0][0]})
    assert r.status_code==409,r.text
    assert r.json()['detail']['code']=='INTEGRATION_EVIDENCE_REVIEW_REQUIRED'
    assert query(ts[0],'SELECT count(*) FROM products WHERE tenant_id=%s')[0][0]==1
