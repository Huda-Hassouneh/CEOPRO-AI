"""Original catalog presentation contracts over canonical Python products/usage.
No separate catalog, inventory or entitlement authority.
"""
import csv,io,json,zipfile
from itertools import islice
from uuid import UUID
from fastapi import APIRouter,Depends,HTTPException,Query,UploadFile,File
from pydantic import BaseModel,Field,ConfigDict
from psycopg2.extras import Json
from src.ai import db
from src.ai.core import usage,text_search
from src.ai.core.barcode import normalize_barcode
from src.ai.core.office_files import ensure_safe_office_archive
from src.ai.core.security import get_tenant_context,TenantContext
from src.ai.routers import products

router=APIRouter(prefix='/catalog/products',tags=['canonical catalog'])

class Version(BaseModel):
    model_config=ConfigDict(extra='forbid')
    expectedVersion:int=Field(gt=0)
class GtinRow(Version):
    productId:UUID
    gtin:str=Field(pattern=r'^(?:[0-9]{8}|[0-9]{12}|[0-9]{13}|[0-9]{14})$')
class GtinBatch(BaseModel):
    model_config=ConfigDict(extra='forbid')
    rows:list[GtinRow]=Field(min_length=1,max_length=500)

def _run(ctx,action,write=False):
    conn=db.app_role_connection(ctx.tenant_id,ctx.user_id)
    try:
        if write:products._require_catalog_admin(conn,ctx.tenant_id,ctx.user_id)
        result=action(conn);conn.commit();return result
    except Exception as exc:
        conn.rollback()
        if getattr(exc,'pgcode',None)=='23505':raise HTTPException(409,'An active product already uses this SKU.') from None
        raise
    finally:conn.close()

def _settings(c,ctx):
    with c.cursor() as q:
        q.execute('SELECT primary_currency,supported_currencies FROM companies WHERE tenant_id=%s',(ctx.tenant_id,));row=q.fetchone()
        primary,supported=row if row else (None,[])
        q.execute("SELECT role_key FROM tenant_users WHERE tenant_id=%s AND user_id=%s AND removed_at IS NULL AND platform_suspended_at IS NULL",(ctx.tenant_id,ctx.user_id));role=q.fetchone()
    return {'primaryCurrency':primary,'currencies':list(dict.fromkeys([x for x in [primary,*(supported or [])] if x])),'canEdit':bool(role and role[0] in ('owner','admin'))}

def _rows(c,tenant,where='TRUE',params=(),suffix=''):
    with c.cursor() as q:
        q.execute("SELECT p.product_id,p.product_name,p.sku,p.metadata->>'barcode',p.current_price,p.cost_price,p.currency,p.brand,p.category,p.description,p.version,p.deleted_at,p.created_at,p.updated_at,p.source FROM products p WHERE p.tenant_id=%s AND "+where+' '+suffix,(tenant,*params))
        return [{'id':str(r[0]),'name':r[1],'sku':r[2],'gtin':r[3],'price':str(r[4]),'costPrice':str(r[5]) if r[5] is not None else None,'currency':r[6],'brand':r[7],'category':r[8],'description':r[9],'version':r[10],'archivedAt':r[11].isoformat() if r[11] else None,'createdAt':r[12].isoformat() if r[12] else None,'updatedAt':r[13].isoformat() if r[13] else None,'source':r[14]} for r in q.fetchall()]

def _duplicates(c,tenant,rows):
    with c.cursor() as q:
        q.execute("SELECT lpad(metadata->>'barcode',14,'0'),count(*) FROM products WHERE tenant_id=%s AND deleted_at IS NULL AND metadata->>'barcode' IS NOT NULL GROUP BY 1",(tenant,));counts=dict(q.fetchall())
    for row in rows:row['gtinDuplicate']=bool(row['gtin'] and counts.get(row['gtin'].zfill(14),0)>1)
    return rows

@router.get('')
def listing(search:str=Query('',max_length=150),status:str=Query('active',pattern='^(active|archived)$'),sort:str=Query('name',pattern='^(name|newest)$'),page:int=Query(1,ge=1),pageSize:int=Query(25,ge=1,le=100),ctx:TenantContext=Depends(get_tenant_context)):
    def action(c):
        where='p.deleted_at IS '+('NULL' if status=='active' else 'NOT NULL');args=[]
        if search:
            where+=" AND (p.product_name::text ILIKE %s OR p.sku ILIKE %s OR p.metadata->>'barcode' ILIKE %s)";args=[text_search.like_pattern(search)]*3
        with c.cursor() as q:
            q.execute('SELECT count(*) FROM products p WHERE p.tenant_id=%s AND '+where,(ctx.tenant_id,*args));total=q.fetchone()[0]
            q.execute('SELECT count(*) FROM products WHERE tenant_id=%s AND deleted_at IS NULL',(ctx.tenant_id,));used=q.fetchone()[0]
        plan=usage.get_plan(c,ctx.tenant_id);feature=plan['features'].get(usage.TRACKED_PRODUCTS);limit=feature['limit'] if feature else 0
        remaining=usage.capacity_left(c,ctx.tenant_id,usage.TRACKED_PRODUCTS)
        order="p.created_at DESC,p.product_id" if sort=='newest' else "COALESCE(p.product_name->>'en',p.product_name->>'ar',''),p.product_id"
        rows=_duplicates(c,ctx.tenant_id,_rows(c,ctx.tenant_id,where,(*args,pageSize,(page-1)*pageSize),'ORDER BY '+order+' LIMIT %s OFFSET %s'))
        return {**_settings(c,ctx),'products':rows,'pagination':{'page':page,'pages':max(1,-(-total//pageSize)),'total':total,'pageSize':pageSize},'capacity':{'currentUsage':used,'limit':limit,'remaining':remaining,'isUnlimited':limit is None}}
    return _run(ctx,action)

@router.get('/{product_id}')
def detail(product_id:UUID,ctx:TenantContext=Depends(get_tenant_context)):
    def action(c):
        rows=_rows(c,ctx.tenant_id,'p.product_id=%s',(str(product_id),))
        if not rows:raise HTTPException(404,'Product not found.')
        product=_duplicates(c,ctx.tenant_id,rows)[0]
        with c.cursor() as q:
            q.execute('SELECT inventory_id,warehouse_location,stock_quantity,updated_at FROM inventory WHERE tenant_id=%s AND product_id=%s ORDER BY updated_at DESC NULLS LAST,inventory_id',(ctx.tenant_id,str(product_id)));levels=q.fetchall()
            q.execute('SELECT price_history_id,old_price,new_price,currency,changed_at FROM product_price_history WHERE tenant_id=%s AND product_id=%s ORDER BY changed_at DESC,price_history_id DESC LIMIT 100',(ctx.tenant_id,str(product_id)));history=q.fetchall()
        return {**_settings(c,ctx),'product':product,'inventory':{'currentStock':float(levels[0][2]) if len(levels)==1 else None,'reason':'ambiguous' if len(levels)>1 else None,'levels':[{'id':str(x[0]),'warehouse':x[1],'quantity':float(x[2]),'updatedAt':x[3].isoformat() if x[3] else None} for x in levels]},'priceHistory':[{'id':str(x[0]),'oldPrice':str(x[1]) if x[1] is not None else None,'oldCurrency':None,'price':str(x[2]),'currency':x[3],'changedAt':x[4].isoformat()} for x in history]}
    return _run(ctx,action)

def _version(c,ctx,pid,version):
    with c.cursor() as q:q.execute('SELECT version,deleted_at FROM products WHERE tenant_id=%s AND product_id=%s FOR UPDATE',(ctx.tenant_id,str(pid)));row=q.fetchone()
    if not row:raise HTTPException(404,'Product not found.')
    if row[0]!=version:raise HTTPException(409,'Product changed. Reload and review again.')
    return row

def _audit(c,ctx,pid,action,values):
    with c.cursor() as q:q.execute("INSERT INTO audit_logs(tenant_id,user_id,action_type,target_table,record_id,changed_data_json) VALUES(%s,%s,%s,'products',%s,%s)",(ctx.tenant_id,ctx.user_id,action,str(pid),Json(values)))

@router.post('/{product_id}/archive')
def archive(product_id:UUID,body:Version,ctx:TenantContext=Depends(get_tenant_context)):
    def action(c):
        row=_version(c,ctx,product_id,body.expectedVersion)
        if row[1] is None:
            products.delete_product(c,ctx.tenant_id,str(product_id));_audit(c,ctx,product_id,'PRODUCT_ARCHIVED',{'previous_version':row[0]})
        return {'ok':True}
    return _run(ctx,action,True)

@router.post('/{product_id}/restore')
def restore(product_id:UUID,body:Version,ctx:TenantContext=Depends(get_tenant_context)):
    def action(c):
        row=_version(c,ctx,product_id,body.expectedVersion)
        if row[1] is not None:
            usage.ensure_capacity(c,ctx.tenant_id,usage.TRACKED_PRODUCTS,1)
            with c.cursor() as q:q.execute('UPDATE products SET deleted_at=NULL,updated_by_user_id=%s WHERE tenant_id=%s AND product_id=%s',(ctx.user_id,ctx.tenant_id,str(product_id)))
            _audit(c,ctx,product_id,'PRODUCT_RESTORED',{'previous_version':row[0]})
        return {'ok':True}
    return _run(ctx,action,True)

def parse_gtins(name,raw):
    try:
        if not 0<len(raw)<=256*1024:raise ValueError('File size')
        if name.lower().endswith('.csv'):
            if b'\0' in raw:raise ValueError('Invalid CSV')
            matrix=list(islice(csv.reader(io.StringIO(raw.decode('utf-8-sig'))),502))
        elif name.lower().endswith('.xlsx'):
            ensure_safe_office_archive(raw)
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                if len(z.infolist())>100 or sum(x.file_size for x in z.infolist())>4*1024*1024:raise ValueError('Archive size')
            from openpyxl import load_workbook
            book=load_workbook(io.BytesIO(raw),read_only=True,data_only=False)
            try:
                if len(book.worksheets)!=1:raise ValueError('One sheet required')
                matrix=[]
                for row in islice(book.active.iter_rows(),502):
                    if any(x.data_type=='f' or (x.value is not None and not isinstance(x.value,str)) for x in row):raise ValueError('Text cells required')
                    matrix.append([x.value or '' for x in row])
            finally:book.close()
        else:raise ValueError('CSV or XLSX required')
        headers=[v.strip().lower() for v in matrix.pop(0)]
        if sorted(headers)!=['gtin','sku'] or not 1<=len(matrix)<=500:raise ValueError('Invalid headers or row count')
        seen=set();result=[]
        for cells in matrix:
            if len(cells)!=2:raise ValueError('Invalid row')
            sku=products.normalize_sku(cells[headers.index('sku')]);gtin=cells[headers.index('gtin')].strip()
            if not sku or sku.lower() in seen or not gtin.isascii() or not gtin.isdecimal() or normalize_barcode(gtin)!=gtin:raise ValueError('Invalid SKU or GTIN')
            seen.add(sku.lower());result.append((sku,gtin))
        return result
    except Exception as exc:raise HTTPException(422,'GTIN_IMPORT_FILE: provide 1?500 distinct SKU/GTIN text rows in a CSV/XLSX up to 256 KB.') from None

@router.post('/gtin-import/preview')
async def preview_gtins(file:UploadFile=File(...),ctx:TenantContext=Depends(get_tenant_context)):
    raw=await file.read(256*1024+1);parsed=parse_gtins(file.filename or '',raw)
    def action(c):
        rows=_duplicates(c,ctx.tenant_id,_rows(c,ctx.tenant_id,'p.deleted_at IS NULL'));by_sku={r['sku'].lower():r for r in rows if r['sku']}
        result=[]
        for sku,gtin in parsed:
            p=by_sku.get(sku.lower())
            if not p:raise HTTPException(422,'GTIN_IMPORT_SKU: no active product matches a file SKU.')
            duplicate=any(r['id']!=p['id'] and r['gtin'] and r['gtin'].zfill(14)==gtin.zfill(14) for r in rows) or sum(v.zfill(14)==gtin.zfill(14) for _,v in parsed)>1
            result.append({'productId':p['id'],'name':p['name'],'sku':p['sku'],'currentGtin':p['gtin'],'gtin':gtin,'expectedVersion':p['version'],'duplicate':duplicate})
        return {'rows':result}
    return _run(ctx,action,True)

@router.post('/gtin-import/apply')
def apply_gtins(body:GtinBatch,ctx:TenantContext=Depends(get_tenant_context)):
    def action(c):
        if len({r.productId for r in body.rows})!=len(body.rows):raise HTTPException(422,'Duplicate product rows.')
        for row in sorted(body.rows,key=lambda r:str(r.productId)):
            current=_version(c,ctx,row.productId,row.expectedVersion)
            if current[1] is not None:raise HTTPException(409,'Product archived. Preview again.')
            try:valid=normalize_barcode(row.gtin)
            except ValueError:raise HTTPException(422,'Invalid GTIN check digit.') from None
            products.update_product(c,ctx.tenant_id,ctx.user_id,str(row.productId),{'barcode':valid})
            _audit(c,ctx,row.productId,'PRODUCT_GTIN_UPDATED',{'gtin':valid,'previous_version':row.expectedVersion})
        return {'updated':len(body.rows)}
    return _run(ctx,action,True)
